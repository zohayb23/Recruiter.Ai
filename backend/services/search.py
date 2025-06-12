from elasticsearch import Elasticsearch
from typing import List, Dict, Any, Optional
import os
import docx
import pdfplumber
import pandas as pd
from pathlib import Path
import logging
import json

logger = logging.getLogger(__name__)

class SearchService:
    def __init__(self):
        self.es = Elasticsearch(os.getenv('ELASTICSEARCH_URL', 'http://localhost:9200'))
        self.index = 'candidates'
        self._ensure_index()
        self._index_resumes()

    def _ensure_index(self):
        """Ensure the index exists with proper mappings"""
        if not self.es.indices.exists(index=self.index):
            mappings = {
                "properties": {
                    "name": {"type": "text"},
                    "skills": {
                        "type": "text",
                        "fields": {
                            "keyword": {"type": "keyword"}
                        }
                    },
                    "experience": {"type": "text"},
                    "education": {"type": "text"},
                    "location": {"type": "keyword"},
                    "current_role": {"type": "text"},
                    "summary": {"type": "text"},
                    "content": {"type": "text"},
                    "filename": {"type": "keyword"},
                    "source": {"type": "keyword"}
                }
            }
            self.es.indices.create(index=self.index, mappings=mappings)

    def _read_resume_content(self, filepath: str) -> str:
        """Read content from different resume file types"""
        try:
            if filepath.endswith('.docx'):
                doc = docx.Document(filepath)
                return "\n".join([paragraph.text for paragraph in doc.paragraphs])
            elif filepath.endswith('.pdf'):
                with pdfplumber.open(filepath) as pdf:
                    return "\n".join([page.extract_text() or "" for page in pdf.pages])
            elif filepath.endswith('.csv'):
                df = pd.read_csv(filepath)
                if 'Resume' in df.columns:
                    return "\n".join(df['Resume'].astype(str).tolist())
            return ""
        except Exception as e:
            print(f"Error reading {filepath}: {str(e)}")
            return ""

    def _index_resumes(self):
        """Index all resumes from different directories"""
        # Index DOCX resumes
        if os.path.exists('docx_resumes'):
            for filename in os.listdir('docx_resumes'):
                if filename.endswith('.docx'):
                    filepath = os.path.join('docx_resumes', filename)
                    content = self._read_resume_content(filepath)
                    if content:
                        self.es.index(
                            index=self.index,
                            document={
                                'content': content,
                                'filename': filename,
                                'source': 'docx'
                            }
                        )

        # Index PDF resumes
        if os.path.exists('pdf_resumes'):
            for filename in os.listdir('pdf_resumes'):
                if filename.endswith('.pdf'):
                    filepath = os.path.join('pdf_resumes', filename)
                    content = self._read_resume_content(filepath)
                    if content:
                        self.es.index(
                            index=self.index,
                            document={
                                'content': content,
                                'filename': filename,
                                'source': 'pdf'
                            }
                        )

        # Index CSV resumes
        if os.path.exists('csv_resumes'):
            for filename in os.listdir('csv_resumes'):
                if filename.endswith('.csv'):
                    filepath = os.path.join('csv_resumes', filename)
                    content = self._read_resume_content(filepath)
                    if content:
                        self.es.index(
                            index=self.index,
                            document={
                                'content': content,
                                'filename': filename,
                                'source': 'csv'
                            }
                        )

    def _build_bool_query(self, query_groups: List[Dict[str, Any]]) -> Dict[str, Any]:
        """Build an Elasticsearch bool query from query groups"""
        bool_query = {
            "bool": {
                "must": [],
                "should": [],
                "must_not": []
            }
        }

        for group in query_groups:
            group_terms = []
            
            for term in group["terms"]:
                term_query = {
                    "multi_match": {
                        "query": term["value"],
                        "fields": ["content^3", "skills^2", "experience", "current_role"],
                        "type": "phrase"
                    }
                }

                if term["operator"] == "NOT":
                    bool_query["bool"]["must_not"].append(term_query)
                else:
                    group_terms.append(term_query)

            if group_terms:
                group_clause = {
                    "bool": {
                        "should" if group["operator"] == "OR" else "must": group_terms,
                        "minimum_should_match": 1 if group["operator"] == "OR" else None
                    }
                }

                if group.get("parentheses"):
                    bool_query["bool"]["must"].append({"bool": {"must": [group_clause]}})
                else:
                    bool_query["bool"]["must"].append(group_clause)

        return bool_query

    async def search_candidates(self, query_groups: List[Dict[str, Any]], page: int = 1, size: int = 20) -> Dict[str, Any]:
        """Search candidates using boolean query groups"""
        try:
            query = self._build_bool_query(query_groups)
            
            response = self.es.search(
                index=self.index,
                query=query,
                from_=(page - 1) * size,
                size=size,
                track_total_hits=True
            )

            return {
                "total": response["hits"]["total"]["value"],
                "results": [
                    {
                        "id": hit["_id"],
                        "score": hit["_score"],
                        "filename": hit["_source"].get("filename", ""),
                        "source": hit["_source"].get("source", ""),
                        "content": hit["_source"].get("content", "")[:500] + "..."  # Return first 500 chars
                    }
                    for hit in response["hits"]["hits"]
                ],
                "page": page,
                "size": size
            }
        except Exception as e:
            print(f"Search error: {str(e)}")
            return {
                "total": 0,
                "results": [],
                "page": page,
                "size": size,
                "error": str(e)
            }

    async def index_candidate(self, candidate: Dict[str, Any]) -> Dict[str, Any]:
        """Index a candidate document"""
        try:
            response = self.es.index(
                index=self.index,
                document=candidate
            )
            return {"success": True, "id": response["_id"]}
        except Exception as e:
            return {"success": False, "error": str(e)}

def read_docx(file_path: str) -> str:
    try:
        doc = docx.Document(file_path)
        return ' '.join([paragraph.text for paragraph in doc.paragraphs])
    except Exception as e:
        logger.error(f"Error reading DOCX file {file_path}: {str(e)}")
        return ""

def read_pdf(file_path: str) -> str:
    try:
        with pdfplumber.open(file_path) as pdf:
            return ' '.join([page.extract_text() or '' for page in pdf.pages])
    except Exception as e:
        logger.error(f"Error reading PDF file {file_path}: {str(e)}")
        return ""

def search_resumes(query: str, search_type: str = "fulltext", top_k: int = 10) -> List[Dict[str, Any]]:
    """
    Search resumes using the specified search type.
    """
    try:
        # Get paths to resume directories
        parent_dir = Path(os.path.dirname(os.path.dirname(os.path.dirname(__file__))))
        pdf_dir = parent_dir / "pdf_resumes"
        docx_dir = parent_dir / "docx_resumes"
        
        logger.info(f"Searching with type {search_type} for query: {query}")
        results = []
        
        # Process each directory
        for directory in [pdf_dir, docx_dir]:
            if directory.exists():
                for file in directory.glob("*.*"):
                    if not file.is_file():
                        continue
                        
                    try:
                        # Read file content based on type
                        content = ""
                        if file.suffix.lower() == '.pdf':
                            content = read_pdf(str(file))
                        elif file.suffix.lower() == '.docx':
                            content = read_docx(str(file))
                        else:
                            continue

                        if not content:
                            logger.warning(f"No content extracted from {file}")
                            continue

                        content_lower = content.lower()
                        query_lower = query.lower()
                        
                        # Calculate match score based on search type
                        match_score = 0
                        if search_type == "fulltext":
                            # Simple keyword matching
                            query_terms = query_lower.split()
                            matches = sum(1 for term in query_terms if term in content_lower)
                            match_score = matches / len(query_terms) if query_terms else 0
                            
                        elif search_type == "semantic":
                            # Basic semantic matching (improve this with actual NLP later)
                            query_terms = set(query_lower.split())
                            content_terms = set(content_lower.split())
                            common_terms = query_terms.intersection(content_terms)
                            match_score = len(common_terms) / len(query_terms) if query_terms else 0
                            
                        elif search_type == "skills":
                            # Skills matching
                            skills = extract_skills(content_lower)
                            query_skills = set(query_lower.split())
                            matching_skills = [skill for skill in skills if any(q in skill.lower() for q in query_skills)]
                            match_score = len(matching_skills) / len(skills) if skills else 0

                        if match_score > 0:
                            # Extract preview and metadata
                            preview = content[:300] + "..." if len(content) > 300 else content
                            metadata = extract_metadata(content)
                            skills = extract_skills(content)

                            result = {
                                "filename": file.name,
                                "content": preview,
                                "match_score": match_score,
                                "skills": skills,
                                "metadata": metadata,
                                "search_type": search_type
                            }
                            results.append(result)
                            logger.info(f"Found match in {file.name} with score {match_score}")
                            
                    except Exception as e:
                        logger.error(f"Error processing file {file}: {str(e)}")
                        continue

        # Sort by score and return top_k results
        results.sort(key=lambda x: x.get("match_score", 0), reverse=True)
        logger.info(f"Returning {len(results[:top_k])} results for {search_type} search")
        return results[:top_k]

    except Exception as e:
        logger.error(f"Search error: {str(e)}")
        return []

def extract_skills(content: str) -> List[str]:
    """Extract skills from resume content."""
    common_skills = [
        "python", "java", "javascript", "typescript", "react", "angular", "vue",
        "node.js", "express", "django", "flask", "fastapi", "spring",
        "sql", "mysql", "postgresql", "mongodb", "redis",
        "aws", "azure", "gcp", "docker", "kubernetes",
        "machine learning", "ai", "data science", "nlp",
        "html", "css", "bootstrap", "tailwind",
        "git", "jenkins", "ci/cd", "agile", "scrum"
    ]
    
    found_skills = []
    content_lower = content.lower()
    for skill in common_skills:
        if skill.lower() in content_lower:
            found_skills.append(skill)
    
    return found_skills

def extract_metadata(content: str) -> Dict[str, Any]:
    """Extract metadata from resume content."""
    metadata = {
        "experience": [],
        "education": []
    }
    
    lines = content.split('\n')
    current_section = None
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        line_lower = line.lower()
        
        # Detect sections
        if "experience" in line_lower or "work history" in line_lower:
            current_section = "experience"
            metadata["experience"].append(line)
        elif "education" in line_lower or "academic" in line_lower:
            current_section = "education"
            metadata["education"].append(line)
        elif current_section and len(line) > 30:  # Only add substantial lines
            metadata[current_section].append(line)
            
    # Limit the number of entries
    metadata["experience"] = metadata["experience"][:5]  # Keep top 5 experiences
    metadata["education"] = metadata["education"][:3]    # Keep top 3 education entries
            
    return metadata 