from typing import List, Dict, Any, Optional
import re
from collections import defaultdict
import os
import docx
import pdfplumber
import pandas as pd
from pathlib import Path

class SkillRatingSystem:
    def __init__(self):
        # Define common skills and their variations
        self.skill_variations = {
            'python': ['python', 'django', 'flask', 'fastapi'],
            'javascript': ['javascript', 'js', 'node.js', 'nodejs', 'react', 'angular'],
            'database': ['sql', 'mysql', 'postgresql', 'mongodb'],
            'cloud': ['aws', 'azure', 'gcp', 'cloud'],
            'devops': ['docker', 'kubernetes', 'jenkins', 'ci/cd'],
        }
    
    def extract_skills(self, text: str) -> Dict[str, float]:
        """Extract skills from text and calculate their frequency."""
        text = text.lower()
        skills = defaultdict(float)
        
        for category, variations in self.skill_variations.items():
            for variation in variations:
                count = len(re.findall(r'\b' + re.escape(variation) + r'\b', text))
                if count > 0:
                    skills[category] += count
        
        return dict(skills)
    
    def calculate_skill_match(self, resume_text: str, required_skills: List[str]) -> Dict[str, float]:
        """Calculate how well a resume matches required skills."""
        resume_skills = self.extract_skills(resume_text)
        skill_scores = {}
        
        for skill in required_skills:
            skill = skill.lower()
            if skill in resume_skills:
                # Direct match
                skill_scores[skill] = min(1.0, resume_skills[skill] / 3)  # Normalize score
            else:
                # Check for related skills
                related_score = 0
                for category, variations in self.skill_variations.items():
                    if skill in variations and category in resume_skills:
                        related_score = min(0.8, resume_skills[category] / 3)  # Slightly lower score for related skills
                skill_scores[skill] = related_score
        
        return skill_scores
    
    def rate_resumes(self, job_description: Optional[str] = None, required_skills: Optional[List[str]] = None) -> List[Dict[str, Any]]:
        """Rate resumes based on job description or required skills."""
        if not job_description and not required_skills:
            raise ValueError("Either job description or required skills must be provided")
        
        # Mock resume data for testing
        resumes = [
            {
                "filename": "john_doe_resume.pdf",
                "source": "pdf",
                "content": "Experienced software engineer with Python, Django, and SQL expertise. Led multiple successful projects using AWS and Docker.",
            },
            {
                "filename": "jane_smith_resume.docx",
                "source": "docx",
                "content": "Full-stack developer proficient in JavaScript, React, and Node.js. Built scalable web applications using MongoDB and Azure.",
            },
            {
                "filename": "bob_wilson_resume.pdf",
                "source": "pdf",
                "content": "Data scientist with expertise in Python, machine learning, and PostgreSQL. Implemented ML models using GCP.",
            }
        ]
        
        # Extract required skills from job description if provided
        if job_description:
            extracted_skills = list(self.extract_skills(job_description).keys())
            required_skills = extracted_skills if extracted_skills else required_skills
        
        if not required_skills:
            return []
        
        # Rate each resume
        rated_results = []
        for resume in resumes:
            skill_scores = self.calculate_skill_match(resume["content"], required_skills)
            
            # Calculate overall score
            overall_score = sum(skill_scores.values()) / len(skill_scores) if skill_scores else 0
            
            rated_results.append({
                "filename": resume["filename"],
                "source": resume["source"],
                "overall_score": round(overall_score * 100, 2),  # Convert to percentage
                "skill_scores": {skill: round(score * 100, 2) for skill, score in skill_scores.items()}
            })
        
        # Sort by overall score
        rated_results.sort(key=lambda x: x["overall_score"], reverse=True)
        return rated_results

def load_resumes() -> List[Dict[str, Any]]:
    """Load resumes from all available sources."""
    resumes = []
    parent_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    
    # Load DOCX resumes
    docx_dir = os.path.join(parent_dir, "..", "docx_resumes")
    if os.path.exists(docx_dir):
        for filename in os.listdir(docx_dir):
            if filename.endswith(".docx"):
                try:
                    doc = docx.Document(os.path.join(docx_dir, filename))
                    content = " ".join([para.text for para in doc.paragraphs])
                    resumes.append({
                        "filename": filename,
                        "source": "docx",
                        "content": content
                    })
                except Exception as e:
                    print(f"Error reading {filename}: {str(e)}")
    
    # Load PDF resumes
    pdf_dir = os.path.join(parent_dir, "..", "pdf_resumes")
    if os.path.exists(pdf_dir):
        for filename in os.listdir(pdf_dir):
            if filename.endswith(".pdf"):
                try:
                    with pdfplumber.open(os.path.join(pdf_dir, filename)) as pdf:
                        content = " ".join([page.extract_text() or "" for page in pdf.pages])
                        resumes.append({
                            "filename": filename,
                            "source": "pdf",
                            "content": content
                        })
                except Exception as e:
                    print(f"Error reading {filename}: {str(e)}")
    
    # Load CSV resumes
    csv_dir = os.path.join(parent_dir, "..", "csv_resumes")
    if os.path.exists(csv_dir):
        for filename in os.listdir(csv_dir):
            if filename.endswith(".csv"):
                try:
                    df = pd.read_csv(os.path.join(csv_dir, filename))
                    if "Resume" in df.columns:
                        for idx, row in df.iterrows():
                            resumes.append({
                                "filename": f"{filename}_{idx}",
                                "source": "csv",
                                "content": str(row["Resume"])
                            })
                except Exception as e:
                    print(f"Error reading {filename}: {str(e)}")
    
    return resumes

def get_skill_variations(skill: str) -> List[str]:
    """Get variations of a skill name."""
    variations = [skill.lower()]
    
    # Add common variations
    if skill.lower() == "python":
        variations.extend(["py", "python3", "python2"])
    elif skill.lower() == "javascript":
        variations.extend(["js", "es6", "ecmascript"])
    elif skill.lower() == "machine learning":
        variations.extend(["ml", "deep learning", "ai", "artificial intelligence"])
    elif skill.lower() == "sql":
        variations.extend(["mysql", "postgresql", "sqlite", "tsql", "plsql"])
    elif skill.lower() == "data analysis":
        variations.extend(["data analytics", "data science", "data mining", "data visualization"])
    
    return variations

def rate_skills(required_skills: List[str], top_k: int = 10) -> List[Dict[str, Any]]:
    """Rate resumes based on required skills."""
    try:
        # Get paths to resume directories
        parent_dir = Path(os.path.dirname(os.path.dirname(os.path.dirname(__file__))))
        pdf_dir = parent_dir / "pdf_resumes"
        docx_dir = parent_dir / "docx_resumes"
        
        results = []
        
        # Process each directory
        for directory in [pdf_dir, docx_dir]:
            if directory.exists():
                for file in directory.glob("*.*"):
                    if not file.is_file():
                        continue
                        
                    try:
                        # Read file content based on type
                        content = ""
                        if file.suffix.lower() == '.pdf':
                            content = read_pdf(str(file))
                        elif file.suffix.lower() == '.docx':
                            content = read_docx(str(file))
                        else:
                            continue

                        if not content:
                            continue

                        # Extract skills from content
                        found_skills = extract_skills(content)
                        
                        # Calculate match score
                        matching_skills = [skill for skill in found_skills if any(req.lower() in skill.lower() for req in required_skills)]
                        score = len(matching_skills) / len(required_skills) if required_skills else 0
                        
                        if score > 0:
                            results.append({
                                "filename": file.name,
                                "content": content,
                                "score": score,
                                "matching_skills": matching_skills,
                                "missing_skills": [skill for skill in required_skills if not any(skill.lower() in found.lower() for found in found_skills)],
                                "source": file.suffix.lower()[1:]  # Remove the dot from extension
                            })
                    except Exception as e:
                        print(f"Error processing {file}: {str(e)}")
                        continue

        # Sort by score and return top results
        results.sort(key=lambda x: x["score"], reverse=True)
        return results[:top_k]

    except Exception as e:
        print(f"Error in rate_skills: {str(e)}")
        return []

def read_pdf(file_path: str) -> str:
    """Read text content from a PDF file."""
    try:
        with pdfplumber.open(file_path) as pdf:
            return " ".join([page.extract_text() or "" for page in pdf.pages])
    except Exception as e:
        print(f"Error reading PDF {file_path}: {str(e)}")
        return ""

def read_docx(file_path: str) -> str:
    """Read text content from a DOCX file."""
    try:
        doc = docx.Document(file_path)
        return " ".join([paragraph.text for paragraph in doc.paragraphs])
    except Exception as e:
        print(f"Error reading DOCX {file_path}: {str(e)}")
        return ""

def extract_skills(text: str) -> List[str]:
    """Extract skills from text using a predefined list."""
    common_skills = [
        'Python', 'Java', 'JavaScript', 'TypeScript', 'React', 'Angular', 'Vue.js',
        'Node.js', 'Express', 'Django', 'Flask', 'Spring', 'AWS', 'Azure', 'GCP',
        'Docker', 'Kubernetes', 'CI/CD', 'Git', 'SQL', 'MongoDB', 'PostgreSQL',
        'Redis', 'GraphQL', 'REST', 'Microservices', 'Machine Learning', 'AI',
        'Data Science', 'TensorFlow', 'PyTorch', 'NLP', 'Agile', 'Scrum',
        'HTML', 'CSS', 'Bootstrap', 'Sass', 'jQuery', 'Redux', 'Next.js',
        'PHP', 'Laravel', 'Ruby', 'Rails', 'C++', 'C#', '.NET', 'Unity',
        'Swift', 'Kotlin', 'Android', 'iOS', 'Mobile Development',
        'DevOps', 'Jenkins', 'Terraform', 'Ansible', 'Linux', 'Unix',
        'Blockchain', 'Ethereum', 'Solidity', 'Smart Contracts',
        'Data Analysis', 'Pandas', 'NumPy', 'Scikit-learn', 'R',
        'UI/UX', 'Figma', 'Adobe XD', 'Photoshop', 'Illustrator'
    ]
    
    found_skills = []
    text_lower = text.lower()
    for skill in common_skills:
        if re.search(r'\b' + re.escape(skill.lower()) + r'\b', text_lower):
            found_skills.append(skill)
    return found_skills 